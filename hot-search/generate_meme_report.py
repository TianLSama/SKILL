# -*- coding: utf-8 -*-
"""网络热梗追踪表生成脚本。

以 example.xlsx 为唯一格式基准，用 openpyxl 显式设置全部样式，
读取 report_data.json（经多源校验的数据）后输出 4 个 Sheet：
总榜 / 分平台 / 潜力榜 / 字段说明与信源。

输出位置：当前项目目录下的 download/ 文件夹（遵守“产物不出项目目录”硬约束）。
"""

import json
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_FILE = os.path.join(BASE_DIR, "report_data.json")
OUT_DIR = os.path.join(BASE_DIR, "download")
OUT_FILE = os.path.join(OUT_DIR, "网络热梗追踪_20260925.xlsx")

DATE_TAG = "20260925"

# ---- 格式常量（与 example.xlsx 逐格一致） ----
FONT_NAME = "微软雅黑"
HEADER_FILL = PatternFill("solid", fgColor="305496")
BLOCK_FILL = PatternFill("solid", fgColor="8EA9DB")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
LINK_FONT_COLOR = "0563C1"
_side = Side(style="thin", color="D9D9D9")
BORDER = Border(left=_side, right=_side, top=_side, bottom=_side)

LIFECYCLE_FILLS = [
    ("🔥", "FCE4D6"),
    ("🚀", "F8CBAD"),
    ("⛰️", "FFF2CC"),
    ("📉", "D9D9D9"),
    ("🔄", "E2EFDA"),
]

HEADERS = [
    "排名", "梗名称", "形式分类", "核心平台", "首发起源", "爆发时间", "生命周期",
    "话题简介", "梗来源/出处", "代表作品标题", "代表作品链接", "话题主页链接",
    "播放量", "点赞数", "评论数", "二创作品数(估)", "相关话题标签", "热度指数",
    "风险提示", "数据抓取时间", "信源链接",
]
WIDTHS = [6, 24, 16, 10, 40, 12, 20, 56, 26, 34, 34, 32, 26, 12, 10, 26, 28, 9, 30, 12, 60]
CENTER_COLS = {0, 2, 3, 5, 17}   # 排名 / 形式分类 / 核心平台 / 爆发时间 / 热度指数
LINK_COLS = {10, 11}             # 代表作品链接 / 话题主页链接
LIFECYCLE_COL = 6

POT_HEADERS = ["序号", "梗名称", "形式", "平台", "话题简介", "代表链接", "数据抓取时间"]
POT_WIDTHS = [6, 26, 18, 10, 60, 44, 14]
POT_CENTER = {0, 2, 3, 6}
POT_LINK = {5}

GROUP_NAMES = ["抖音区（含双平台）", "B站区", "全网区"]


def lifecycle_fill(value):
    if not value:
        return None
    for keyword, color in LIFECYCLE_FILLS:
        if keyword in value:
            return PatternFill("solid", fgColor=color)
    return None


def header_cell(cell):
    cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER


def data_cell(cell, col_idx, center_cols, link_cols, lifecycle_col=None, borders=True):
    cell.font = Font(name=FONT_NAME, size=10)
    if col_idx in center_cols:
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    else:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    if borders:
        cell.border = BORDER
    if lifecycle_col is not None and col_idx == lifecycle_col:
        fill = lifecycle_fill(cell.value)
        if fill is not None:
            cell.fill = fill
    if col_idx in link_cols and isinstance(cell.value, str) and cell.value.startswith("http"):
        cell.hyperlink = cell.value
        cell.font = Font(name=FONT_NAME, size=10, color=LINK_FONT_COLOR, underline="single")


def write_header(ws, row, headers):
    for i, text in enumerate(headers, start=1):
        c = ws.cell(row, i, text)
        header_cell(c)
    ws.row_dimensions[row].height = 28


def write_total_rows(ws, start_row, rows):
    r = start_row
    for row in rows:
        for i, value in enumerate(row, start=1):
            c = ws.cell(r, i, value)
            data_cell(c, i - 1, CENTER_COLS, LINK_COLS, LIFECYCLE_COL)
        r += 1
    return r


def set_widths(ws, widths):
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def build_zongbang(wb, total):
    ws = wb.create_sheet("总榜")
    set_widths(ws, WIDTHS)
    write_header(ws, 1, HEADERS)
    # 按热度指数降序重排并重新编号（就地更新 ordered，供「分平台」沿用全局排名）
    ordered = [list(r) for r in total]
    ordered.sort(key=lambda r: (r[17] if isinstance(r[17], (int, float)) else 0), reverse=True)
    for idx, row in enumerate(ordered, start=1):
        row[0] = idx
        for i, value in enumerate(row, start=1):
            c = ws.cell(idx + 1, i, value)
            data_cell(c, i - 1, CENTER_COLS, LINK_COLS, LIFECYCLE_COL)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:U{len(ordered) + 1}"
    return ordered


def block_of(platform, name):
    if name == "抖音区（含双平台）":
        return platform in ("抖音", "双平台")
    if name == "B站区":
        return platform == "B站"
    return platform == "全网"


def build_fenpingtai(wb, ordered):
    ws = wb.create_sheet("分平台")
    set_widths(ws, WIDTHS)
    row = 1
    for gi, name in enumerate(GROUP_NAMES):
        rows = [r for r in ordered if block_of(r[3], name)]
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=21)
        title = ws.cell(row, 1, f"{name}（{len(rows)}条）")
        title.font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
        title.fill = BLOCK_FILL
        title.alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 22
        row += 1
        write_header(ws, row, HEADERS)
        row += 1
        for data_row in rows:
            for i, value in enumerate(data_row, start=1):
                c = ws.cell(row, i, value)
                data_cell(c, i - 1, CENTER_COLS, LINK_COLS, LIFECYCLE_COL)
            row += 1
        if gi < len(GROUP_NAMES) - 1:
            row += 1  # 空行分隔
    return ws


def build_potential(wb, potential):
    ws = wb.create_sheet("潜力榜")
    set_widths(ws, POT_WIDTHS)
    write_header(ws, 1, POT_HEADERS)
    r = 2
    for idx, row in enumerate(potential, start=1):
        row = list(row)
        row[0] = idx
        for i, value in enumerate(row, start=1):
            c = ws.cell(r, i, value)
            data_cell(c, i - 1, POT_CENTER, POT_LINK)
        r += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{r - 1}"
    return ws


def build_meta(wb, meta):
    ws = wb.create_sheet("字段说明与信源")
    set_widths(ws, [20, 110])
    for r, pair in enumerate(meta, start=1):
        a_val = pair[0] if len(pair) > 0 else None
        b_val = pair[1] if len(pair) > 1 else None
        a = ws.cell(r, 1, a_val)
        b = ws.cell(r, 2, b_val)
        a.font = Font(name=FONT_NAME, size=10, bold=True)
        b.font = Font(name=FONT_NAME, size=10)
        a.alignment = Alignment(vertical="top", wrap_text=True)
        b.alignment = Alignment(vertical="top", wrap_text=True)
        if isinstance(a_val, str) and a_val.startswith("【"):
            a.fill = NOTE_FILL
    return ws


def main():
    with open(DATA_FILE, "r", encoding="utf-8") as f:
        data = json.load(f)
    total = data["total"]
    potential = data["potential"]
    meta = data["meta"]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ordered = build_zongbang(wb, total)
    build_fenpingtai(wb, ordered)
    build_potential(wb, potential)
    build_meta(wb, meta)

    os.makedirs(OUT_DIR, exist_ok=True)
    wb.save(OUT_FILE)
    sys.stdout.write(
        f"saved: {OUT_FILE}\n"
        f"总榜 {len(ordered)} 条 | 潜力榜 {len(potential)} 条 | meta {len(meta)} 行\n"
    )


if __name__ == "__main__":
    main()
